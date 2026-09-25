#!/usr/bin/env python3
"""Google Local Services Ads (LSA) advertiser list builder.

Finds home-service companies that are CURRENTLY paying for Google Local Services Ads
("Google Guaranteed" listings) in a set of cities, records the proof, then finds each
company's real phone line, website and owner so they can be cold-called.

Why this works without an API: Google's "More providers" page
(google.com/localservices/prolist?q=<trade city state>) lists only paying LSA advertisers.
Every card carries Google's own data-test-id="paid-list-card" marker and the advertiser's
Google Ads customer id. We save the rendered page, a screenshot and the card data as evidence.

Steps (each writes to out/ and can be re-run; nothing is re-fetched that is already saved):

  lists      load every (city x trade x query variant) listing page -> out/cards.jsonl + evidence
  profiles   open the LSA profile of the top N advertisers -> weekly hours, license, ownership notes
  enrich     real phone + website + owner from the company website, our BBB list, web search
  export     dedupe, rank, split across 3 callers -> out/lsa_leads.csv, out/caller_N.csv, out/evidence.md
  run        all four

  python3 lsa_leads.py run --cities Tulsa --top 20            # first-20 check for one city
  python3 lsa_leads.py run                                     # all cities, all trades
  python3 lsa_leads.py export --top 500

IMPORTANT: the phone number shown inside a Local Services ad is a Google call-tracking number.
Calling it creates a billable lead for the business. This tool records it only as evidence
(column lsa_display_phone, DO NOT CALL) and puts the company's own line in the Phone column.

Requires: pip install playwright requests beautifulsoup4 lxml ddgs && playwright install chromium
"""
import argparse
import base64
import csv
import hashlib
import html as htmlmod
import json
import math
import os
import random
import re
import sys
import time
import urllib.parse
from datetime import datetime, timezone
from pathlib import Path

import requests
from bs4 import BeautifulSoup

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36")

# city, state, suburbs that share the metro (LSA results are location sensitive, suburbs surface more advertisers)
CITIES = [
    ("Tulsa", "OK", ["Broken Arrow", "Owasso", "Bixby"]),
    ("Omaha", "NE", ["Bellevue", "Papillion", "Council Bluffs"]),
    ("Wichita", "KS", ["Derby", "Andover", "Haysville"]),
    ("Sioux Falls", "SD", ["Brandon", "Harrisburg", "Tea"]),
    ("Lincoln", "NE", ["Waverly", "Hickman", "Seward"]),
    ("Springfield", "MO", ["Nixa", "Ozark", "Republic"]),
    ("Des Moines", "IA", ["West Des Moines", "Ankeny", "Urbandale"]),
    ("Grand Rapids", "MI", ["Wyoming", "Kentwood", "Grandville"]),
    ("Lexington", "KY", ["Nicholasville", "Georgetown", "Richmond"]),
    ("Fort Wayne", "IN", ["New Haven", "Huntertown", "Auburn"]),
]
STATE_FOR_SUBURB = {"Council Bluffs": "IA"}
TRADES = {  # first term is the primary one, also used for suburb queries
    "Roofing": ["roofers", "roofing contractors", "roof repair"],
    "HVAC": ["hvac", "air conditioning repair", "furnace repair"],
    "Plumbing": ["plumbers", "drain cleaning", "water heater repair"],
    "Electrical": ["electricians", "electrical contractors", "electrical repair"],
}
CORPORATE_REMOVE = [  # corporate-owned chains and private-equity platforms: removed from the list
    "service experts", "ars ", "ars/", "rescue rooter", "roto-rooter", "roto rooter", "erie home", "erie metal",
    "power home", "leaffilter", "leaf filter", "leafguard", "window world", "champion window", "bath fitter", "re-bath",
    "sears", "lowe's", "home depot", "homeserve", "apex service partners", "wrench group", "goettl", "any hour", "anyhour",
    "redwood services", "turnpoint", "horizon services", "morris-jenkins", "ace hardware", "ace home services",
    "hometown services", "legacy service partners", "southern hvac", "heartland home", "frontdoor",
    "american residential", "comfort systems usa", "1-800", "1800", "hvac.com", "sila heating", "sila services",
    "one hour air conditioning & heating of", "unique indoor comfort",
    "homex", "nexstar", "arco ", "arcoair", "abc home", "abc plumbing", "four seasons heating", "dabella",
    "renewal by andersen", "west shore home", "long home", "windows usa", "1-800-hansons", "hansons", "feldco",
    "pella", "andersen", "power home", "champion", "leafguard", "thompson creek", "mad city", "k designers",
    "mr. roof", "mr roof", "bone dry roofing", "atrium home", "heartland home services", "northwinds", "best choice roofing", "peterman brothers",
]
FRANCHISE = [  # franchise brands: kept when the location is locally owned, but flagged and penalised
    "mr. rooter", "mr rooter", "zoom drain", "rooter-man", "rooterman", "bluefrog", "benjamin franklin", "one hour",
    "mister sparky", "mr. electric", "mr electric", "aire serv", "aireserv", "neighborly", "precision door",
    "dr. roof", "roof maxx", "roofmaxx", "mighty dog roofing", "handyman connection", "mr. handyman", "1 tom plumber",
    "pink plumber", "the plumbing pros", "hoffmann brothers", "wire wiz", "electricians on call", "lightning bug",
    "captain electric", "storm guard", "roof squad", "moxie", "gold medal", "grounds guys", "wireman", "temperaturepro",
    "temperature pro",
]
ACQUIRED_RE = re.compile(r"(was |been |were )?acquired by|acquisition by|private[- ]equity|sold to |owned by [^;]*(group|partners|"
                         r"holdings|capital)|part of the [^;]* group|portfolio company|backed by", re.I)
BUTTONS = {"Get phone number", "Book", "Get quote", "Share", "Book online", "Call", "Message", "Website"}
BLOCK_DOMAINS = (
    "google.", "facebook.", "instagram.", "yelp.", "bbb.org", "yellowpages", "mapquest", "indeed.", "glassdoor",
    "angi.", "homeadvisor", "thumbtack", "houzz", "x.com", "twitter.", "youtube.", "tiktok.", "wikipedia",
    "manta.", "dnb.com", "buzzfile", "chamberofcommerce", "birdeye", "nextdoor", "porch.", "bark.com", "projectmapit", "igolocal", "plumberssupply",
    "bing.", "duckduckgo", "pinterest", "reddit.", "opencorporates", "bizapedia", "buildzoom", "superpages",
    "citysearch", "foursquare", "trustpilot", "linkedin.", "zoominfo", "crunchbase", "apple.com", "yellowbook",
    "expertise.com", "threebestrated", "hometown", "nicelocal", "cylex", "brownbook", "localsearch", "findglocal",
    "roofingcompanies", "networx", "porch", "homeguide", "servicewhale", "yelp", "houzz",
    "chamber", "electricianslist", "plumberslist", "rooferslist", "serviceagent", "poyst", "mepo.org", "lovable",
    ".gov", "homestars", "elocal", "hotfrog", "cybo", "n49.", "2findlocal", "merchantcircle", "ezlocal", "storeboard",
    "local.com", "yellowbot", "insiderpages", "citysquares", "showmelocal", "ebusinesspages", "tupalo", "spoke.com",
    "lead411", "rocketreach", "signalhire", "dexknows", "iglobal", "localstack", "findacontractor", "contractors.com",
    "hvacrepair", "yellowpagecity", "usaroofers", "roofingcontractor.", "mapsconnect", "cityfos", "bizhwy", "salespider",
    "companieshub", "callupcontact", "trustanalytica", "prolocalservices", "top10", "bestprosintown", "hubbiz", "loc8nearme",
    "hvacservice.io", "plumbingservice.io", "roofingservice.io", "electricservice.io", "servicepros", "homeflow", "nearby",
    "trane.com", "lennox.com", "carrier.com", "americanstandardair", "rheem.com", "goodman", "bryant.com", "daikin",
    "guildquality", "voolt", "tydl.io", "searchaplumber", "hvacloc", "prosforhome", "omaha-ne.com", "gaf.com", "owenscorning",
    "certainteed", "servicequoteai", "reputation.", "topservdigital", "bwpsites", "memberzone", "patch.com", "nextdoor",
    "alignable", "yellowpages", "bizapedia", "chamberofcommerce", "valuenews", "kompass", "localz", "wheree", "homeyou",
)
REGIONAL_CHAIN = ["paschal", "shamrock roofing", "dabella"]   # multi-state operators: branch can't buy locally
GENERIC_TOKENS = {"roofing", "roofers", "roofer", "roof", "roofs", "plumbing", "plumber", "plumbers", "electric", "electrical",
                  "electricians", "electrician", "hvac", "heating", "cooling", "air", "conditioning", "service", "services",
                  "home", "homes", "pros", "solutions", "company", "companies", "contractors", "contractor", "construction",
                  "exteriors", "drain", "sewer", "septic", "mechanical", "repair", "repairs", "restoration", "gutters",
                  "siding", "tulsa", "omaha", "wichita", "sioux", "falls", "lincoln", "springfield", "moines", "grand",
                  "rapids", "lexington", "fort", "wayne", "oklahoma", "kansas", "nebraska", "iowa", "michigan", "kentucky",
                  "indiana", "missouri", "dakota", "quality", "best", "pro", "american", "family", "city", "metro"}


def site_matches(name, url, title, allow_domain=True):
    """A found website counts only if its domain or page title carries a distinctive word of the business name."""
    words = norm_name(name).split()
    toks = [t for t in words if len(t) >= 4 and t not in GENERIC_TOKENS]
    dom = urllib.parse.urlsplit(url).netloc.lower().replace("-", "")
    if any(b.replace("-", "") in dom for b in BLOCK_DOMAINS):
        return False                                   # known directory / dealer-locator / social domains
    if re.search(r"people|news|magazine|times|journal|tribune|gazette|herald|world|press|directory|guide|review|rated|"
                 r"nearme|pages|listing|listings|citation|profile|wiki|blog|forum|jobs|career|indeed|hba|association|members|chamber", dom):
        return False                                   # media / directory sites, even when they mention the name
    compact = "".join(words)
    if allow_domain and compact and len(compact) >= 6 and compact in dom:
        return True                                    # abestroofing.com, americanhomepros.com
    if not toks:  # names made only of generic words: require the first two words together in the title
        first2 = " ".join(words[:2])
        return bool(first2) and first2 in norm_name(title or "")
    if allow_domain and any(t in dom for t in toks):
        return True
    if not allow_domain and not re.search(r"roof|plumb|electric|hvac|heat|air|cool|drain|sewer|contract|service|mechanical",
                                          (title or "").lower()):
        return False                                   # a guessed domain must at least look like a trade business
    tl = " " + re.sub(r"[^a-z0-9 ]", " ", (title or "").lower()) + " "
    return sum(1 for t in toks if f" {t} " in tl or t in tl.replace(" ", "")) >= min(2, len(toks))
DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
AREA_CODES = {  # a phone outside the state's area codes gets a "verify" note (could be a national office)
    "OK": {"918", "539", "405", "572", "580"}, "NE": {"402", "531", "308"}, "KS": {"316", "785", "913", "620"},
    "SD": {"605"}, "MO": {"417", "573", "314", "636", "816", "660", "557"}, "IA": {"515", "319", "563", "641", "712"},
    "MI": {"616", "231", "269", "517", "313", "248", "734", "810", "586", "989", "906", "947", "679"},
    "KY": {"859", "502", "606", "270", "364"}, "IN": {"260", "317", "463", "219", "574", "765", "812", "930"},
}


def log(msg):
    try:
        print(msg, flush=True)
    except BrokenPipeError:
        sys.exit(0)


def now_iso():
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def digits(s):
    d = re.sub(r"\D", "", s or "")
    return d[1:] if len(d) == 11 and d.startswith("1") else d


def fmt_phone(d):
    d = digits(d)
    return f"({d[:3]}) {d[3:6]}-{d[6:]}" if len(d) == 10 else ""


def norm_name(n):
    n = htmlmod.unescape(n or "").lower()
    n = re.split(r"\s[|\-–]\s", n)[0]                    # "Half Moon Plumbing | Tulsa" -> "half moon plumbing"
    n = re.sub(r"[^a-z0-9 ]+", " ", n)
    n = re.sub(r"\b(llc|inc|co|corp|corporation|company|the|of|and|ltd|l l c)\b", " ", n)
    return re.sub(r"\s+", " ", n).strip()


def slug(s):
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


def read_jsonl(p):
    if not Path(p).exists():
        return []
    return [json.loads(l) for l in open(p, encoding="utf-8") if l.strip()]


def append_jsonl(p, rows):
    with open(p, "a", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")


def write_csv(path, rows, cols):
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in cols})


# ----------------------------------------------------------------------------------------------
# browser
# ----------------------------------------------------------------------------------------------
class Blocked(Exception):
    pass


class Browser:
    """One headless Chromium for the whole run. Google shows a CAPTCHA page ("/sorry/") when it
    rate-limits; we detect it, wait, and retry once."""

    def __init__(self, headless=True, min_wait=4.0):
        from playwright.sync_api import sync_playwright
        self._pw = sync_playwright().start()
        kw = {"headless": headless, "args": ["--no-sandbox", "--disable-quic", "--disable-blink-features=AutomationControlled"]}
        exe = os.environ.get("CHROMIUM_PATH")
        if exe:
            kw["executable_path"] = exe
        px = os.environ.get("HTTPS_PROXY") or os.environ.get("https_proxy")
        if px:
            kw["proxy"] = {"server": px}
        self.browser = self._pw.chromium.launch(**kw)
        self.ctx = self.browser.new_context(locale="en-US", viewport={"width": 1280, "height": 2200}, user_agent=UA)
        self.page = self.ctx.new_page()
        self.min_wait = min_wait
        self.loads = 0

    def get(self, url, screenshot=None, settle=5.0):
        for attempt in range(2):
            self.loads += 1
            r = self.page.goto(url, wait_until="domcontentloaded", timeout=90000)
            time.sleep(settle + random.uniform(0, 1.5))
            final = self.page.url
            html = self.page.content()
            if "/sorry/" in final or "unusual traffic" in html[:20000] or (r and r.status == 429):
                if attempt == 0:
                    log("    Google is rate-limiting this network (CAPTCHA page). Waiting 10 minutes before one retry...")
                    time.sleep(600)
                    continue
                raise Blocked(url)
            if screenshot:
                try:
                    self.page.screenshot(path=str(screenshot), full_page=True)
                except Exception:
                    pass
            time.sleep(self.min_wait + random.uniform(0, 3))
            return r.status if r else 0, final, html
        raise Blocked(url)

    def close(self):
        try:
            self.browser.close()
            self._pw.stop()
        except Exception:
            pass


# ----------------------------------------------------------------------------------------------
# listing pages
# ----------------------------------------------------------------------------------------------
def build_queries(cities, trades):
    qs = []
    for city, state, suburbs in cities:
        for trade, terms in trades.items():
            for t in terms:
                qs.append({"city": city, "state": state, "metro": city, "trade": trade, "term": t, "q": f"{t} {city} {state}"})
            for sub in suburbs:
                st = STATE_FOR_SUBURB.get(sub, state)
                qs.append({"city": sub, "state": st, "metro": city, "trade": trade, "term": terms[0], "q": f"{terms[0]} {sub} {st}"})
    return qs


def prolist_url(q):
    return "https://www.google.com/localservices/prolist?ssta=1&q=" + urllib.parse.quote(q) + "&src=1"


def _af_blocks(html):
    out = {}
    for b in re.findall(r"AF_initDataCallback\((\{.*?\})\);", html, re.S):
        k = re.search(r"key: '([^']+)'", b)
        m = re.search(r"data:(.*), sideChannel", b, re.S)
        if k and m:
            try:
                out[k.group(1)] = json.loads(m.group(1))
            except Exception:
                pass
    return out


def _walk(o, path=()):
    if isinstance(o, list):
        for i, v in enumerate(o):
            yield from _walk(v, path + (i,))
    elif isinstance(o, dict):
        for k, v in o.items():
            yield from _walk(v, path + (k,))
    else:
        yield path, o


def _json_extras(html):
    """From the page's embedded data: per advertiser name -> ad display phone, booking link."""
    extras = {}
    d = _af_blocks(html).get("ds:1")
    if d is None:
        return extras
    for path, v in _walk(d):
        if isinstance(v, str) and v.startswith("https://www.google.com/localservices/profile") and path and path[-1] == 0:
            try:
                rec_list = d
                for k in path[:-2]:
                    rec_list = rec_list[k]
                name_list = rec_list[path[-2]]
                name = name_list[1] if len(name_list) > 1 and isinstance(name_list[1], str) else ""
            except Exception:
                continue
            if not name:
                continue
            e = extras.setdefault(name, {})
            for p2, v2 in _walk(rec_list):
                if isinstance(v2, str):
                    if re.fullmatch(r"\+1\d{10}", v2) and "lsa_display_phone" not in e:
                        e["lsa_display_phone"] = fmt_phone(v2)
                    elif v2.startswith("/localservices/booking?ebd=") and "booking_url" not in e:
                        try:
                            raw = base64.urlsafe_b64decode(v2.split("ebd=")[1] + "==")
                            m = re.search(rb"https?://[A-Za-z0-9._\-/]+", raw)
                            if m:
                                e["booking_url"] = m.group(0).decode()[:160]
                        except Exception:
                            pass
    return extras


def parse_cards(html, meta, evidence_png):
    soup = BeautifulSoup(html, "lxml")
    cards = soup.select('div[data-test-id="paid-list-card"]') or soup.select('div[jscontroller="xkZ6Lb"]')
    extras = _json_extras(html)
    rows = []
    for c in cards:
        lines = [l.strip() for l in c.get_text("\n", strip=True).split("\n") if l.strip()]
        if not lines:
            continue
        row = {"name": lines[0], "rating": "", "reviews": 0, "category": "", "years": "", "serves": "",
               "hours_status": "", "highlights": [], "reply": "", "has_booking": "Book" in lines or "Book online" in lines,
               "customer_id": c.get("data-customer-id", ""), "test_id": c.get("data-test-id", ""),
               "preview_ad": c.get("data-is-preview-ad", ""),
               "profile_url": "https://www.google.com" + htmlmod.unescape(c.get("data-profile-url-path", "")),
               "query": meta["q"], "query_url": prolist_url(meta["q"]), "metro": meta["metro"], "city": meta["city"],
               "state": meta["state"], "trade": meta["trade"], "term": meta["term"], "seen_at": now_iso(),
               "evidence_png": str(evidence_png)}
        i = 1
        if i < len(lines) and re.fullmatch(r"\d\.\d", lines[i]):
            row["rating"] = lines[i]; i += 1
        if i < len(lines) and re.fullmatch(r"\((\d[\d,]*)\)", lines[i]):
            row["reviews"] = int(lines[i].strip("()").replace(",", "")); i += 1
        if i < len(lines) and lines[i] not in BUTTONS and not re.search(r"years?|Serves|Open|Closed|Opened", lines[i]):
            row["category"] = lines[i]; i += 1
        hours = []
        for l in lines[i:]:
            if l in BUTTONS:
                continue
            m = re.match(r"(\d+)\+?\s+years? in business", l)
            if m:
                row["years"] = int(m.group(1)); continue
            m = re.match(r"Opened in (\d{4})", l)
            if m:
                row["years"] = max(0, datetime.now().year - int(m.group(1))); continue
            if l == "1 year in business":
                row["years"] = 1; continue
            m = re.match(r"Serves (.+)", l)
            if m:
                row["serves"] = m.group(1); continue
            if re.match(r"^(Open|Closed)\b", l) or l.startswith("·"):
                hours.append(l.replace(" ", " ")); continue
            m = re.match(r"Typically replies in (.+)", l)
            if m:
                row["reply"] = m.group(1); continue
            row["highlights"].append(l)
        row["hours_status"] = " ".join(hours).replace("Open ·", "Open ·").strip()
        row["highlights"] = " · ".join(h for h in row["highlights"] if h)
        row.update(extras.get(row["name"], {}))
        rows.append(row)
    return rows


def cmd_lists(a):
    out = Path(a.out); raw = out / "raw"; ev = out / "evidence"
    raw.mkdir(parents=True, exist_ok=True); ev.mkdir(parents=True, exist_ok=True)
    cities = [c for c in CITIES if not a.cities or c[0].lower() in [x.lower() for x in a.cities]]
    trades = {k: v for k, v in TRADES.items() if not a.trades or k.lower() in [x.lower() for x in a.trades]}
    queries = build_queries(cities, trades)
    if a.limit:
        queries = queries[: a.limit]
    done = {r["query"] for r in read_jsonl(out / "cards.jsonl")}
    todo = [q for q in queries if q["q"] not in done]
    log(f"lists: {len(queries)} queries, {len(todo)} still to load")
    if not todo:
        return
    br = Browser(headless=not a.headed, min_wait=a.wait)
    try:
        for i, q in enumerate(todo, 1):
            s = slug(q["q"])
            png = ev / f"{s}.png"
            try:
                status, final, html = br.get(prolist_url(q["q"]), screenshot=png)
            except Blocked:
                log("lists: blocked twice, stopping. Re-run later; everything so far is saved.")
                break
            (raw / f"{s}.html").write_text(html, encoding="utf-8")
            rows = parse_cards(html, q, png.relative_to(out))
            if not rows:  # record the query as done even when Google shows no advertisers here
                rows = [{"name": "", "query": q["q"], "query_url": prolist_url(q["q"]), "metro": q["metro"], "city": q["city"],
                         "state": q["state"], "trade": q["trade"], "term": q["term"], "seen_at": now_iso(), "customer_id": "",
                         "empty": True}]
            append_jsonl(out / "cards.jsonl", rows)
            n = sum(1 for r in rows if r.get("name"))
            log(f"  [{i}/{len(todo)}] {n:2d} advertisers  <- {q['q']}")
    finally:
        br.close()
    allrows = [r for r in read_jsonl(out / "cards.jsonl") if r.get("name")]
    log(f"lists: {len(allrows)} card appearances, {len({r['customer_id'] or r['name'] for r in allrows})} unique advertisers -> {out / 'cards.jsonl'}")


# ----------------------------------------------------------------------------------------------
# merge / score (used by profiles, enrich, export)
# ----------------------------------------------------------------------------------------------
def merge_cards(out, cities=None):
    biz = {}
    want = {c.lower() for c in (cities or [])}
    for r in read_jsonl(Path(out) / "cards.jsonl"):
        if not r.get("name") or (want and r.get("metro", "").lower() not in want):
            continue
        key = r.get("customer_id") or norm_name(r["name"])
        b = biz.get(key)
        if b is None:
            b = dict(r); b["trades"] = set(); b["queries"] = []; b["cities_seen"] = set(); b["metros"] = set()
            b["first_seen"] = r["seen_at"]
            biz[key] = b
        b["trades"].add(r["trade"]); b["queries"].append(r["query"]); b["cities_seen"].add(r["city"]); b["metros"].add(r["metro"])
        if (r.get("reviews") or 0) > (b.get("reviews") or 0):
            b["reviews"], b["rating"] = r["reviews"], r["rating"]
        for k in ("lsa_display_phone", "booking_url", "years", "hours_status", "highlights", "reply"):
            if r.get(k) and not b.get(k):
                b[k] = r[k]
    # same company advertising under two LSA accounts (one per license): merge when one normalized
    # name is a word-prefix of the other within the same metro
    keys = sorted(biz, key=lambda k: -int(biz[k].get("reviews") or 0))
    merged_into = {}
    for i, k1 in enumerate(keys):
        if k1 in merged_into:
            continue
        n1 = norm_name(biz[k1]["name"]).split()
        for k2 in keys[i + 1:]:
            if k2 in merged_into or biz[k1]["metro"] != biz[k2]["metro"]:
                continue
            n2 = norm_name(biz[k2]["name"]).split()
            short, long_ = (n1, n2) if len(n1) <= len(n2) else (n2, n1)
            if len(short) >= 2 and long_[:len(short)] == short:
                a_, b_ = biz[k1], biz[k2]
                a_["trades"] |= b_["trades"]; a_["queries"] += b_["queries"]; a_["cities_seen"] |= b_["cities_seen"]
                a_.setdefault("other_customer_ids", []).append(b_["customer_id"])
                for f in ("lsa_display_phone", "booking_url", "years", "hours_status", "highlights", "reply"):
                    if b_.get(f) and not a_.get(f):
                        a_[f] = b_[f]
                merged_into[k2] = k1
    for k in merged_into:
        del biz[k]
    for b in biz.values():
        b["trades"] = sorted(b["trades"]); b["cities_seen"] = sorted(b["cities_seen"]); b["metros"] = sorted(b["metros"])
        b["primary_metro"] = b["metro"]
    return biz


def is_franchise(name):
    n = " " + name.lower() + " "
    return any(f and f in n for f in FRANCHISE)


def is_corporate(name):
    n = " " + name.lower() + " "
    return any(f and f in n for f in CORPORATE_REMOVE)


OWNER_TITLES = re.compile(r"owner|founder|president|principal|proprietor|general manager|managing", re.I)


def score(b):
    """Ranking + the two explanation columns. Sets b['score'], b['why'], b['flags'], b['exclude'].
    Ideal lead: independent, established, 150-3,000 reviews, owner still involved, 24/7 phone-heavy."""
    why, flags = [], []
    reviews = int(b.get("reviews") or 0)
    s = math.log10(reviews + 1) * 8                          # 150 reviews ~ 17, 1,000 ~ 24, 5,000 ~ 30
    if 150 <= reviews <= 3000:
        s += 6; why.append(f"{reviews:,} reviews: real call volume, still owner-sized")
    elif reviews < 150:
        s -= (150 - reviews) / 150 * 6
        if reviews < 60:
            flags.append(f"only {reviews} reviews: may not have much inbound volume")
    elif reviews <= 5000:
        s -= 3; why.append(f"{reviews:,} reviews: high volume")
    elif reviews <= 8000:
        s -= 9; flags.append(f"{reviews:,} reviews: looks large, probably has a dispatch desk or call center")
    else:
        s -= 15; flags.append(f"{reviews:,} reviews: very large operation, likely corporate-style call center")
    try:
        rt = float(b.get("rating") or 0)
    except ValueError:
        rt = 0
    if rt >= 4.7:
        s += 4; why.append(f"rated {rt}")
    elif rt >= 4.5:
        s += 2
    elif rt and rt < 4.3:
        s -= 3; flags.append(f"rating {rt}")
    yrs = b.get("years")
    if isinstance(yrs, int):
        s += min(yrs, 25) / 25 * 6
        if yrs >= 10:
            why.append(f"{yrs}+ years in business")
        elif yrs < 3:
            s -= 2; flags.append(f"{yrs} years in business")
    hs = (b.get("hours_status") or "").lower()
    if "24 hours" in hs:
        s += 6; why.append("advertises 24-hour phones: every night call is exposure")
    elif re.search(r"closes (7|8|9|10|11)", hs):
        s += 4; why.append("open late")
    elif hs:
        s += 1
    nt = len(b.get("trades", []))
    if nt == 2:
        s += 2; why.append("advertises in 2 trades")
    elif nt >= 3:
        s -= 3; flags.append("advertises in 3+ trades: multi-service operation")
    nq = len(set(b.get("queries", [])))
    if nq >= 4:
        s += 3; why.append(f"shows in {nq} different LSA searches: broad ad coverage")
    if b.get("booking_url"):
        s += 2; why.append("online booking in the ad")
    rp = (b.get("reply") or "").lower()
    if "min" in rp:
        s += 2; why.append(f"replies in {b['reply']}")
    elif "hour" in rp:
        s += 1
    hl = (b.get("highlights") or "").lower()
    if "bbb accredited" in hl:
        s += 2
    if re.search(r"family owned|veteran owned|local business|locally", hl + " " + " ".join(b.get("ownership", [])).lower()):
        s += 3; why.append("family / locally owned")
    e = b.get("e", {})
    if e.get("website"):
        s += 3
        sig = e.get("site_signals", "")
        if re.search(r"Google Ads tag|call tracking", sig):
            s += 2; why.append("website runs Google Ads tag / call tracking: pays for calls beyond LSA")
        if "answering service" in sig or "Smith.ai" in sig:
            s -= 4; flags.append("website mentions an answering service: may already have coverage")
        if "Housecall Pro" in sig or "ServiceTitan" in sig or "Jobber" in sig:
            s += 1; why.append("runs field-service software")
    else:
        s -= 2; flags.append("no verified website found")
    own = e.get("owner", "")
    if own:
        if OWNER_TITLES.search(e.get("title", "") or "owner"):
            s += 7; why.append(f"owner-level contact known: {own}")
        else:
            s += 3; why.append(f"contact known: {own} ({e.get('title', '')})")
    else:
        flags.append("owner not identified yet: ask for the owner")
    if e.get("phone"):
        s += 4
    else:
        s -= 8; flags.append("no verified business line: look up before calling")
    if is_franchise(b["name"]):
        s -= 8; flags.append("franchise brand: confirm the local owner can buy")
    n = " " + b["name"].lower() + " "
    if any(x and x in n for x in REGIONAL_CHAIN) or len(b.get("metros", [])) >= 2:
        s -= 8; flags.append("multi-city / multi-state operator: the local office may not make the buying decision")
    exclude = ""
    if is_corporate(b["name"]):
        exclude = "corporate-owned chain / national operator"
    note_text = " ".join([e.get("notes", ""), b.get("override_note", "")])
    if ACQUIRED_RE.search(note_text):
        exclude = exclude or "acquired / private-equity owned: decision not local"
    if b.get("exclude_override"):
        exclude = b["exclude_override"]
    if re.search(r"area code \d{3} is not local", e.get("notes", "")):
        flags.append("phone area code not local: verify")
    s += b.get("adjust", 0)
    b["score"] = round(s, 1)
    b["why"] = "; ".join(why)
    b["flags"] = "; ".join(flags)
    b["exclude"] = exclude
    return s


# ----------------------------------------------------------------------------------------------
# profiles
# ----------------------------------------------------------------------------------------------
def parse_profile(html):
    d = _af_blocks(html)
    info = {"weekly_hours": {}, "license": "", "ownership": [], "service_area": []}
    for key in ("ds:3", "ds:1", "ds:2", "ds:4"):
        blk = d.get(key)
        if blk is None:
            continue
        for path, v in _walk(blk):
            if isinstance(v, str):
                if v in DAYS and path:
                    try:
                        rec = blk
                        for k in path[:-1]:
                            rec = rec[k]
                        hrs = [x for p, x in _walk(rec[3]) if isinstance(x, str)] if len(rec) > 3 else []
                        if hrs:
                            info["weekly_hours"][v] = " / ".join(hrs)
                    except Exception:
                        pass
                elif re.search(r"(Registration|License|Lic\.?)\s*#?\s*[A-Z0-9-]{4,}", v) and len(v) < 80:
                    info["license"] = v
                elif re.search(r"(?i)(locally|family|veteran|women|minority)[- ]owned", v) and len(v) < 60:
                    if v not in info["ownership"]:
                        info["ownership"].append(v)
    return info


def cmd_profiles(a):
    out = Path(a.out); raw = out / "raw_profiles"; raw.mkdir(parents=True, exist_ok=True)
    biz = merge_cards(out, a.cities)
    have = {p["customer_id"] for p in read_jsonl(out / "profiles.jsonl")}
    for b in biz.values():
        score(b)
    todo = sorted([b for b in biz.values() if b["customer_id"] and b["customer_id"] not in have], key=lambda b: -b["score"])
    todo = todo[: a.top]
    log(f"profiles: {len(todo)} profile pages to load ({len(have)} done)")
    if not todo:
        return
    br = Browser(headless=not a.headed, min_wait=a.wait)
    try:
        for i, b in enumerate(todo, 1):
            try:
                status, final, html = br.get(b["profile_url"], settle=4.0)
            except Blocked:
                log("profiles: blocked, stopping; re-run later")
                break
            (raw / f"{b['customer_id']}.html").write_text(html, encoding="utf-8")
            info = parse_profile(html)
            info["customer_id"] = b["customer_id"]; info["name"] = b["name"]; info["fetched_at"] = now_iso()
            append_jsonl(out / "profiles.jsonl", [info])
            wh = info["weekly_hours"]
            log(f"  [{i}/{len(todo)}] {b['name'][:34]:34s} hours: {len(wh)} days  {info['license'][:30]:30s} {' '.join(info['ownership'])[:40]}")
    finally:
        br.close()


# ----------------------------------------------------------------------------------------------
# enrich: real phone, website, owner
# ----------------------------------------------------------------------------------------------
class Web:
    def __init__(self, out):
        self.cache = Path(out) / "cache"; self.cache.mkdir(parents=True, exist_ok=True)
        self.s = requests.Session(); self.s.headers.update({"User-Agent": UA, "Accept-Language": "en-US,en;q=0.9"})
        self._ddgs = None

    def search(self, q, n=8):
        key = hashlib.md5(q.encode()).hexdigest(); cf = self.cache / f"s_{key}.json"
        if cf.exists():
            return json.loads(cf.read_text())
        if self._ddgs is None:
            from ddgs import DDGS
            self._ddgs = DDGS()
        for backend in ("auto", "bing"):
            try:
                res = self._ddgs.text(q, max_results=n, backend=backend)
                out = [{"title": r.get("title", ""), "href": r.get("href", ""), "body": r.get("body", "")} for r in res]
                if out:
                    cf.write_text(json.dumps(out)); time.sleep(2 + random.random() * 2)
                    return out
            except Exception as e:
                if "no results" not in str(e).lower():
                    time.sleep(3)
            time.sleep(1)
        return []

    def get(self, url, timeout=20):
        try:
            r = self.s.get(url, timeout=timeout, allow_redirects=True)
            return r if r.status_code == 200 else None
        except requests.RequestException:
            return None


ROLE = r"(?:Owner|Co-Owner|Owner/Operator|Founder|Co-Founder|President|CEO|General Manager|Managing Partner|Owner & Operator)"
NAME = r"([A-Z][a-z]+(?:\s(?:[A-Z]\.|[A-Z][a-z]+|[A-Z][a-z]+-[A-Z][a-z]+)){1,2})"
OWNER_PATS = [
    re.compile(NAME + r"\s*[,|\-–—:]\s*" + ROLE + r"\b"),
    re.compile(r"\b" + ROLE + r"\s*[,|\-–—:]\s*" + NAME),
    re.compile(r"\b(?i:founded|started|established|owned and operated|owned & operated|is owned|run)\s+(?i:by)\s+" + NAME),
    re.compile(r"\b" + ROLE + r"\s+" + NAME + r"\b"),
]
NAME_STOP = {"Google Guaranteed", "Free Estimate", "Read More", "Contact Us", "Learn More", "Call Now", "Our Team",
             "Better Business", "Family Owned", "Locally Owned", "Home Services", "Air Conditioning", "Roofing Company",
             "Get Started", "Privacy Policy", "Since", "Serving", "Est", "Meet The", "About Us", "Heating Cooling"}


LEAD_WORDS = {"Contact", "Call", "Meet", "Ask", "Email", "Text", "Welcome", "Thanks", "Thank", "Hi", "Hello", "About",
              "Message", "Visit", "Dear", "From", "With", "Our", "The", "By", "Owner", "President", "Founder", "Mr", "Mrs", "Ms", "Dr"}


def find_owner(text):
    text = re.sub(r"\s+", " ", htmlmod.unescape(text))
    for pat in OWNER_PATS:
        for m in pat.finditer(text):
            words = m.group(1).strip().split()
            while words and words[0].rstrip(".") in LEAD_WORDS:
                words = words[1:]
            name = " ".join(words[:3])
            if name in NAME_STOP or len(words) < 2:
                continue
            if re.search(r"\b(Roofing|Plumbing|Electric|HVAC|Heating|Cooling|Services?|Company|Inc|LLC|Home|Air)\b", name):
                continue
            ctx = text[max(0, m.start() - 60): m.end() + 60]
            return name, ctx.strip()
    return "", ""


def site_contact(web, url):
    r = web.get(url)
    if not r:
        return {}
    html_text = r.text
    soup = BeautifulSoup(html_text, "lxml")
    info = {"website": f"{urllib.parse.urlsplit(r.url).scheme}://{urllib.parse.urlsplit(r.url).netloc}"}
    m = re.search(r'href="tel:([^"]+)"', html_text, re.I)
    if m and len(digits(m.group(1))) == 10:
        info["phone"], info["phone_source"] = fmt_phone(m.group(1)), "website tel: link"
    else:
        txt = soup.get_text(" ", strip=True)
        m = re.search(r"(?<!\d)\(?([2-9]\d{2})\)?[\s.-]?(\d{3})[\s.-]?(\d{4})(?!\d)", txt)
        if m:
            info["phone"], info["phone_source"] = f"({m.group(1)}) {m.group(2)}-{m.group(3)}", "website text"
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    info["site_title"] = title[:100]
    # tech signals
    low = html_text.lower()
    sig = []
    for k, label in (("callrail", "CallRail call tracking"), ("calltrackingmetrics", "CTM call tracking"), ("googleadservices", "Google Ads tag"),
                     ("gtag('config', 'aw-", "Google Ads tag"), ("facebook.com/tr", "Meta pixel"), ("housecallpro", "Housecall Pro"),
                     ("servicetitan", "ServiceTitan"), ("getjobber", "Jobber"), ("podium", "Podium chat"), ("smith.ai", "Smith.ai"),
                     ("24/7", "says 24/7"), ("answering service", "mentions answering service"), ("live chat", "live chat")):
        if k in low and label not in sig:
            sig.append(label)
    info["site_signals"] = ", ".join(sig)
    owner, ctx = find_owner(soup.get_text(" ", strip=True))
    if owner:
        info["owner"], info["owner_source"], info["owner_context"] = owner, r.url, ctx[:160]
    else:  # try an about/team page linked from the homepage
        for a_ in soup.find_all("a", href=True):
            h = a_["href"]
            if re.search(r"about|team|our-story|who-we-are|meet", h, re.I) and not h.startswith(("mailto:", "tel:", "#")):
                u2 = urllib.parse.urljoin(r.url, h)
                if urllib.parse.urlsplit(u2).netloc != urllib.parse.urlsplit(r.url).netloc:
                    continue
                r2 = web.get(u2)
                if r2:
                    owner, ctx = find_owner(BeautifulSoup(r2.text, "lxml").get_text(" ", strip=True))
                    if owner:
                        info["owner"], info["owner_source"], info["owner_context"] = owner, r2.url, ctx[:160]
                        break
                break
    return info


def guess_site(web, name, trades, city):
    """Try obvious domains (acmeroofing.com, acmeroofingtulsa.com ...) and keep one whose page mentions the name."""
    base = re.sub(r"[^a-z0-9]", "", norm_name(name))
    words = norm_name(name).split()
    first = "".join(words[:2]) if len(words) > 2 else base
    tw = {"Roofing": "roofing", "HVAC": "hvac", "Plumbing": "plumbing", "Electrical": "electric"}
    stems = [base, first, base.replace("and", "")] + [first + tw[t] for t in trades if t in tw] + \
            [base + re.sub(r"[^a-z]", "", city.lower())]
    cands = []
    for stem in dict.fromkeys(stems):
        if 4 <= len(stem) <= 40:
            cands += [f"https://www.{stem}.com", f"https://{stem}.com"]
    for u in cands[:10]:
        r = web.get(u, timeout=10)
        if not r:
            continue
        m = re.search(r"<title[^>]*>([^<]{0,200})", r.text, re.I)
        title = htmlmod.unescape(m.group(1)) if m else ""
        if site_matches(name, r.url, title, allow_domain=False):   # a guessed domain must prove itself by its title
            return r.url
    return ""


def bbb_search(web, name, city, state, _retry=False):
    """BBB's search page embeds each result as JSON (name, address, main phone, rating, accreditation,
    categories, service area). Profile pages sit behind a bot check, the search page does not."""
    q = urllib.parse.urlencode({"find_country": "USA", "find_text": name, "find_loc": f"{city}, {state}"})
    key = hashlib.md5(("bbb|" + q).encode()).hexdigest(); cf = web.cache / f"b_{key}.html"
    if cf.exists():
        html_text = cf.read_text(encoding="utf-8")
    else:
        r = web.get("https://www.bbb.org/search?" + q, timeout=30)
        if not r:
            return {}
        html_text = r.text; cf.write_text(html_text, encoding="utf-8"); time.sleep(1.5 + random.random())
    want = norm_name(name).split()
    best = None
    for m in re.finditer(r'"businessName":"(.*?)","address":"(.*?)","city":"(.*?)","state":"(.*?)","postalcode":"(.*?)","tobText":"(.*?)","rating":"(.*?)","bbbMember":(true|false)', html_text):
        bname = htmlmod.unescape(m.group(1).replace("<em>", "").replace("</em>", ""))
        got = norm_name(bname).split()
        short, long_ = (want, got) if len(want) <= len(got) else (got, want)
        if not short or long_[:len(short)] != short or m.group(4).upper() != state.upper():
            continue
        tail = html_text[m.end(): m.end() + 3000]
        phones = re.findall(r'\(\d{3}\) \d{3}-\d{4}', tail.split('"location"')[0]) if '"phone":' in tail[:1200] else []
        url = re.search(r'"reportUrl":"([^"]+)"', tail)
        oob = re.search(r'"outOfBusinessStatus":(null|"[^"]*")', tail)
        cand = {"bbb_name": bname, "bbb_city": m.group(3), "bbb_zip": m.group(5), "bbb_rating": m.group(7),
                "bbb_accredited": m.group(8) == "true", "bbb_phone": phones[0] if phones else "",
                "bbb_url": "https://www.bbb.org" + url.group(1) if url else "", "bbb_address": m.group(2),
                "bbb_out_of_business": bool(oob and oob.group(1) != "null")}
        if cand["bbb_out_of_business"]:
            continue
        metro_cities = {city.lower()} | {sub.lower() for c, st, subs in CITIES if c.lower() == city.lower() for sub in subs}
        in_metro = cand["bbb_city"].lower() in metro_cities
        if best is None or (in_metro and best["bbb_city"].lower() not in metro_cities):
            best = cand
    if best is None and len(want) > 2 and not _retry:
        return bbb_search(web, " ".join(want[:2]), city, state, _retry=True)   # "Mullin Plumbing, HVAC & Septic" -> "mullin plumbing"
    return best or {}


def load_bbb_index():
    """Owner names, phones and websites from the BBB-built lists already in this repo."""
    idx_name, idx_phone = {}, {}
    for fn in ("roofing/leads.js", "pest/leads.js"):
        p = REPO / fn
        if not p.exists():
            continue
        s = p.read_text(encoding="utf-8")
        try:
            leads = json.loads(s[s.index("window.LEADS = ") + 15:].strip().rstrip(";"))
        except Exception:
            continue
        for l in leads:
            idx_name[(norm_name(l.get("business", "")), (l.get("state") or "").upper())] = l
            if digits(l.get("phone", "")):
                idx_phone[digits(l["phone"])] = l
    return idx_name, idx_phone


def cmd_fixsites(a):
    """For businesses whose override row names a website, fetch it and record the site phone / signals
    (enriched_zz_manual.jsonl, read last so it wins) when the row has no phone yet."""
    out = Path(a.out)
    biz = merge_cards(out, a.cities)
    have = read_enriched(out)
    op = out / "owner_overrides.csv"
    if not op.exists():
        op = HERE / "data" / "owner_overrides.csv"
    overrides = {}
    with open(op, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            overrides[(norm_name(r.get("Business", "")), (r.get("City") or "").lower())] = r
    web = Web(out)
    target = out / "enriched_zz_manual.jsonl"
    done = {e["customer_id"] for e in read_jsonl(target)} if target.exists() else set()
    n = 0
    for b in biz.values():
        ov = overrides.get((norm_name(b["name"]), b["primary_metro"].lower()))
        if not ov or not (ov.get("Website") or "").strip() or b["customer_id"] in done:
            continue
        e = dict(have.get(b["customer_id"], {}))
        w = ov["Website"].strip()
        if not w.startswith("http"):
            w = "https://" + w
        same = e.get("website", "").rstrip("/").lower() == w.rstrip("/").lower()
        if e.get("phone") and (same or not e.get("phone_source", "").startswith("website")):
            continue                                            # already has a phone from BBB or this very site
        info = site_contact(web, w)
        row = {"customer_id": b["customer_id"], "name": b["name"], "website": w, "website_source": "hand-verified",
               "site_title": info.get("site_title", ""), "site_signals": info.get("site_signals", "")}
        if info.get("phone"):
            row["phone"], row["phone_source"] = info["phone"], info["phone_source"] + " (hand-verified site)"
        for k in ("bbb", "bbb_address", "years_bbb", "notes"):
            if e.get(k):
                row[k] = e[k]
        with open(target, "a", encoding="utf-8") as f:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")
        n += 1
        log(f"  [{n}] {b['name']} ({b['primary_metro']}): {w} -> {row.get('phone', 'no phone on site')}")
    log(f"fixsites: {n} sites fetched")


def read_enriched(out):
    """All enrichment rows across shards (enriched.jsonl, enriched_<shard>.jsonl); later rows win."""
    rows = {}
    for f in sorted(Path(out).glob("enriched*.jsonl")):
        if "_v" in f.name:
            continue
        for e in read_jsonl(f):
            old = rows.get(e["customer_id"])
            if old is None or e.get("phone") or not old.get("phone"):
                rows[e["customer_id"]] = e
    return rows


def cmd_enrich(a):
    out = Path(a.out)
    biz = merge_cards(out, a.cities)
    for b in biz.values():
        score(b)
    have = read_enriched(out)
    target = out / (f"enriched_{a.shard}.jsonl" if a.shard else "enriched.jsonl")
    done = {cid for cid, e in have.items() if e.get("phone") or a.no_retry}   # rows without a phone get another try
    todo = sorted([b for b in biz.values() if b["customer_id"] not in done], key=lambda b: -b["score"])[: a.top]
    log(f"enrich: {len(todo)} businesses ({len(have)} done)")
    web = Web(out)
    bbb_name, bbb_phone = load_bbb_index()
    for i, b in enumerate(todo, 1):
        e = {"customer_id": b["customer_id"], "name": b["name"], "phone": "", "phone_source": "", "website": "", "owner": "",
             "title": "", "owner_source": "", "owner_context": "", "site_signals": "", "bbb": "", "enriched_at": now_iso()}
        # 1. our own BBB list (owner first name, phone, website)
        l = bbb_name.get((norm_name(b["name"]), b["state"]))
        if l:
            e["bbb"] = f"BBB {l.get('rating', '')}{' accredited' if l.get('accredited') else ''}".strip()
            if l.get("owner"):
                e["owner"], e["title"], e["owner_source"] = l["owner"], l.get("title") or "Owner", "BBB list (first name)"
            if l.get("website"):
                e["website"] = l["website"]
            if l.get("phone"):
                e["phone"], e["phone_source"] = l["phone"], "BBB list"
        # 2. BBB search page: main phone, address, rating, accreditation (all trades, no bot check)
        bb = bbb_search(web, b["name"], b["city"], b["state"])
        if bb:
            e["bbb"] = (e["bbb"] or f"BBB {bb['bbb_rating']}{' accredited' if bb['bbb_accredited'] else ''}").strip()
            e["bbb_url"] = bb["bbb_url"]; e["address"] = f"{bb['bbb_address']}, {bb['bbb_city']}, {b['state']} {bb['bbb_zip']}"
            if bb["bbb_phone"] and not e["phone"]:
                e["phone"], e["phone_source"] = bb["bbb_phone"], "BBB listing"
        # 3. website: from BBB row, else guess the domain (cheap), else a web search (slow, rate-limited)
        site = e["website"] or guess_site(web, b["name"], b["trades"], b["city"])
        if not site and not a.no_search:
            res = web.search(f'"{b["name"]}" {b["city"]} {b["state"]}', 8)
            for r in res:
                dom = urllib.parse.urlsplit(r["href"]).netloc.lower()
                if dom and not any(x in dom for x in BLOCK_DOMAINS):
                    site = r["href"]; break
        if site:
            info = site_contact(web, site)
            if info and not site_matches(b["name"], info.get("website", ""), info.get("site_title", "")) \
                    and e.get("phone_source") != "BBB list":
                e["notes"] = f"unverified site candidate {info.get('website', '')} ignored"
                info = {}
                g = guess_site(web, b["name"], b["trades"], b["city"])
                if g:
                    info = site_contact(web, g)
            if info:
                e["website"] = info.get("website", e["website"])
                e["site_title"] = info.get("site_title", "")
                e["site_signals"] = info.get("site_signals", "")
                if info.get("phone") and not e["phone"]:
                    e["phone"], e["phone_source"] = info["phone"], info["phone_source"]
                elif info.get("phone") and e["phone"] and digits(info["phone"]) != digits(e["phone"]):
                    e["phone_alt"] = info["phone"] + " (website)"
                elif info.get("phone") and e["phone"] and digits(info["phone"]) == digits(e["phone"]):
                    e["phone_source"] += " + website"
                if info.get("owner") and (not e["owner"] or len(e["owner"].split()) == 1):
                    e["owner"], e["owner_source"], e["owner_context"] = info["owner"], info["owner_source"], info.get("owner_context", "")
                    e["title"] = e["title"] or (re.search(ROLE, info.get("owner_context", "")) or [None])[0] or "Owner"
                    if isinstance(e["title"], re.Match):
                        e["title"] = e["title"].group(0)
        ac = digits(e["phone"])[:3] if e["phone"] else ""
        if ac and ac not in AREA_CODES.get(b["state"], set()):
            e["notes"] = (e.get("notes", "") + f"; area code {ac} is not local to {b['state']}: verify before calling").strip("; ")
        # never use the ad's tracking number as the phone; note if the site's number equals it
        if b.get("lsa_display_phone") and e["phone"] and digits(b["lsa_display_phone"]) == digits(e["phone"]):
            e["phone_source"] += " (same number as in the ad)"
        append_jsonl(target, [e])
        log(f"  [{i}/{len(todo)}] {b['name'][:32]:32s} | {e['phone'] or '-':14s} {e['phone_source'][:22]:22s} | {e['owner'][:20]:20s} | {urllib.parse.urlsplit(e['website']).netloc[:28] if e['website'] else '-':28s} | {e.get('bbb', '')}")


# ----------------------------------------------------------------------------------------------
# export
# ----------------------------------------------------------------------------------------------
COLS = ["Rank", "List", "Caller", "Top 50", "Business", "Trade", "City", "State", "Phone", "Phone source", "Website",
        "Owner", "Title", "Owner source", "Google rating", "Review count", "Years in business", "Hours (ad status)",
        "Weekly hours", "Why this lead ranks highly", "Red flags", "Highlights", "License", "Ownership notes",
        "Booking tool", "Site signals", "Address (BBB)", "BBB", "Notes",
        "LSA proof: profile URL", "LSA proof: Google Ads customer ID", "LSA proof: listing URL", "LSA proof: screenshot",
        "First seen (UTC)", "Also listed for", "LSA display phone (DO NOT CALL - tracking number)", "Score"]
REMOVED_COLS = ["Business", "Trade", "City", "State", "Google rating", "Review count", "Reason removed",
                "LSA proof: Google Ads customer ID", "LSA proof: profile URL"]


def build_rows(a):
    """Merge cards + profiles + enrichment + overrides, score, dedupe, return (kept_rows, removed_rows)."""
    out = Path(a.out)
    biz = merge_cards(out, a.cities)
    prof = {p["customer_id"]: p for p in read_jsonl(out / "profiles.jsonl")}
    enr = read_enriched(out)
    for b in biz.values():                      # a merged twin's data counts for the survivor
        for cid in b.get("other_customer_ids", []):
            if cid in prof and b["customer_id"] not in prof:
                prof[b["customer_id"]] = prof[cid]
            if cid in enr and (b["customer_id"] not in enr or not enr[b["customer_id"]].get("phone")):
                enr[b["customer_id"]] = enr[cid]
    overrides = {}
    op = out / "owner_overrides.csv"
    if not op.exists():
        op = HERE / "data" / "owner_overrides.csv"     # the hand-verified file committed with the tool
    if op.exists():
        with open(op, encoding="utf-8-sig", newline="") as f:
            for r in csv.DictReader(f):
                overrides[(norm_name(r.get("Business", "")), (r.get("City") or "").lower())] = r
    for key, b in biz.items():
        p = prof.get(b["customer_id"], {}); e = dict(enr.get(b["customer_id"], {}))
        ov = overrides.get((norm_name(b["name"]), b["primary_metro"].lower()))
        if ov:
            if (ov.get("Owner") or "").strip() == "-":            # hand-checked: no owner name is published; drop any scraped guess
                e["owner"], e["title"], e["owner_source"] = "", "", ""
            elif ov.get("Owner"):
                e["owner"] = ov["Owner"]; e["title"] = ov.get("Title", "") or e.get("title", "")
                e["owner_source"] = ov.get("Source", "hand-verified")
            if (ov.get("Website") or "").strip():                  # hand-verified website wins over the scraped candidate
                w = ov["Website"].strip()
                if not w.startswith("http"):
                    w = "https://" + w
                if e.get("website") and e["website"].rstrip("/").lower() != w.rstrip("/").lower() and e.get("phone_source", "").startswith("website"):
                    e["phone"], e["phone_source"] = "", ""      # the phone came from the wrong site
                e["website"], e["site_title"], e["website_source"] = w, b["name"], "hand-verified"
            if (ov.get("Phone") or "").strip():
                e["phone"], e["phone_source"] = fmt_phone(ov["Phone"]), "hand-verified (" + (ov.get("Source") or "web search") + ")"
            if ov.get("Note"):
                b["override_note"] = ov["Note"]
            if ov.get("Exclude"):
                b["exclude_override"] = ov["Exclude"]
            try:
                b["adjust"] = float(ov.get("Adjust") or 0)
            except ValueError:
                b["adjust"] = 0
        if e.get("website") and e.get("website_source") != "hand-verified" and not site_matches(b["name"], e["website"], e.get("site_title", "")):
            e["notes"] = (e.get("notes", "") + f"; website candidate {e['website']} not verified").strip("; ")
            e["website"] = ""
        if e.get("owner"):
            m = re.match(r"^(.*?)[\s,]+(Owner|President|Founder|CEO|Co-Owner|General Manager)$", e["owner"].strip(), re.I)
            if m:
                e["owner"], e["title"] = m.group(1).strip(" ,"), (e.get("title") or m.group(2))
        # a website-sourced phone with a non-local area code usually means a same-named company elsewhere
        ac = digits(e.get("phone", ""))[:3]
        if ac and ac not in AREA_CODES.get(b["state"], set()) and e.get("phone_source", "").startswith("website"):
            e["notes"] = (e.get("notes", "") + f"; dropped {e['website']} / {e['phone']}: same-named company in another state").strip("; ")
            e["phone"], e["phone_source"], e["website"] = "", "", ""
            if e.get("owner_source", "").startswith("http"):
                e["owner"], e["title"], e["owner_source"] = "", "", ""
        b["weekly_hours"] = p.get("weekly_hours", {}); b["ownership"] = p.get("ownership", [])
        b["p"] = p; b["e"] = e
        score(b)
    dom_count = {}
    for b in biz.values():
        w = b["e"].get("website", "")
        if w:
            dom_count[urllib.parse.urlsplit(w).netloc.lower()] = dom_count.get(urllib.parse.urlsplit(w).netloc.lower(), 0) + 1
    for b in biz.values():
        w = b["e"].get("website", "")
        if w and dom_count.get(urllib.parse.urlsplit(w).netloc.lower(), 0) > 1:
            b["e"]["notes"] = (b["e"].get("notes", "") + f"; {w} is shared by several businesses (directory), ignored").strip("; ")
            b["e"]["website"] = ""
            if b["e"].get("phone_source", "").startswith("website"):
                b["e"]["phone"], b["e"]["phone_source"] = "", ""
            score(b)
    ordered = sorted(biz.values(), key=lambda b: -b["score"])
    kept, removed, seen_phone = [], [], {}
    for b in ordered:
        e, p = b["e"], b["p"]
        if b["exclude"]:
            removed.append({"Business": b["name"], "Trade": " / ".join(b["trades"]), "City": b["primary_metro"], "State": b["state"],
                            "Google rating": b.get("rating", ""), "Review count": b.get("reviews", 0), "Reason removed": b["exclude"],
                            "LSA proof: Google Ads customer ID": b["customer_id"], "LSA proof: profile URL": b["profile_url"]})
            continue
        ph = digits(e.get("phone", ""))
        if ph and ph in seen_phone:      # same line under two names: keep the higher-ranked one, note the alias
            seen_phone[ph]["Notes"] += f"; also advertises as {b['name']}"
            continue
        wk = p.get("weekly_hours", {})
        weekly = "; ".join(f"{d[:3]} {wk[d]}" for d in DAYS if d in wk)
        notes = [e.get("bbb", ""), e.get("notes", ""), b.get("override_note", "")]
        if b.get("reply"):
            notes.append(f"replies in {b['reply']}")
        if e.get("phone_alt"):
            notes.append(f"alt phone {e['phone_alt']}")
        row = {
            "Rank": 0, "List": "", "Caller": "", "Top 50": "", "Business": b["name"], "Trade": " / ".join(b["trades"]),
            "City": b["primary_metro"], "State": b["state"], "Phone": e.get("phone", ""), "Phone source": e.get("phone_source", ""),
            "Website": e.get("website", ""), "Owner": e.get("owner", ""), "Title": e.get("title", ""),
            "Owner source": e.get("owner_source", ""), "Google rating": b.get("rating", ""), "Review count": b.get("reviews", 0),
            "Years in business": b.get("years", ""), "Hours (ad status)": b.get("hours_status", ""), "Weekly hours": weekly,
            "Why this lead ranks highly": b["why"], "Red flags": b["flags"], "Highlights": b.get("highlights", ""),
            "License": p.get("license", ""), "Ownership notes": " · ".join(p.get("ownership", [])),
            "Booking tool": urllib.parse.urlsplit(b.get("booking_url", "")).netloc if b.get("booking_url") else "",
            "Site signals": e.get("site_signals", ""), "Address (BBB)": e.get("address", ""), "BBB": e.get("bbb", ""),
            "Notes": "; ".join(n for n in notes if n),
            "LSA proof: profile URL": b["profile_url"], "LSA proof: Google Ads customer ID": b["customer_id"]
            + ("" if not b.get("other_customer_ids") else " (+" + ", ".join(b["other_customer_ids"]) + ")"),
            "LSA proof: listing URL": b["query_url"], "LSA proof: screenshot": b.get("evidence_png", ""),
            "First seen (UTC)": b["first_seen"], "Also listed for": "; ".join(sorted(set(b["queries"]))),
            "LSA display phone (DO NOT CALL - tracking number)": b.get("lsa_display_phone", ""), "Score": b["score"],
            "_reviews": int(b.get("reviews") or 0), "_owner_ok": bool(e.get("owner")) and bool(OWNER_TITLES.search(e.get("title", "") or "owner")),
            "_24h": "24 hours" in (b.get("hours_status") or "").lower(), "_flagged": bool(b["flags"] and re.search(r"franchise|answering service|very large|not local", b["flags"])),
        }
        kept.append(row)
        if ph:
            seen_phone[ph] = row
    return kept, removed


def cmd_export(a):
    out = Path(a.out)
    kept, removed = build_rows(a)
    total = a.top or 500
    primary_n = min(a.primary, total)
    rows = kept[:total]
    for i, r in enumerate(rows, 1):
        r["Rank"] = i
        r["List"] = "Primary" if i <= primary_n else "Backup"
        r["Caller"] = (i - 1) % 3 + 1 if i <= primary_n else ""
    # Top 50: strongest mix of LSA spend, missed-call exposure and reachable ownership
    def top50_key(r):
        k = r["Score"]
        k += 6 if r["_owner_ok"] else -10
        k += 4 if 150 <= r["_reviews"] <= 3000 else 0
        k += 3 if r["_24h"] else 0
        k -= 8 if r["_flagged"] else 0
        k -= 6 if not r["Phone"] else 0
        return -k
    for r in sorted(rows, key=top50_key)[:50]:
        r["Top 50"] = "yes"
    write_csv(out / "lsa_leads.csv", rows, COLS)
    write_csv(out / "primary_300.csv", [r for r in rows if r["List"] == "Primary"], COLS)
    write_csv(out / "backup_200.csv", [r for r in rows if r["List"] == "Backup"], COLS)
    for c in (1, 2, 3):
        write_csv(out / f"caller_{c}.csv", [r for r in rows if r["Caller"] == c], COLS)
    write_csv(out / "top_50.csv", [r for r in rows if r["Top 50"]], COLS)
    write_csv(out / "removed.csv", removed, REMOVED_COLS)
    write_csv(out / "all_ranked.csv", kept, COLS)
    with open(out / "evidence.md", "w", encoding="utf-8") as f:
        f.write(f"# LSA advertiser evidence ({len(rows)} businesses, generated {now_iso()})\n\n")
        f.write("Every row was captured from google.com/localservices/prolist, the page that lists only paying Local Services Ads "
                "advertisers. Each card carried Google's `data-test-id=\"paid-list-card\"` marker and the advertiser's Google Ads "
                "customer id. Listing URL + screenshot reproduce the capture.\n\n")
        f.write("| # | Business | Trade | City | Rating (reviews) | Ads customer ID | Seen (UTC) | Listing query | Profile |\n|---|---|---|---|---|---|---|---|---|\n")
        for r in rows:
            f.write(f"| {r['Rank']} | {r['Business']} | {r['Trade']} | {r['City']}, {r['State']} | {r['Google rating']} ({r['Review count']}) | "
                    f"{r['LSA proof: Google Ads customer ID']} | {r['First seen (UTC)']} | [{r['Also listed for'].split(';')[0]}]({r['LSA proof: listing URL']}) | "
                    f"[profile]({r['LSA proof: profile URL']}) |\n")
    write_workbook(out / "lsa_leads.xlsx", rows, removed, kept)
    with_phone = sum(1 for r in rows if r["Phone"]); with_owner = sum(1 for r in rows if r["Owner"])
    log(f"export: {len(kept)} eligible after removing {len(removed)}; delivered {len(rows)} "
        f"({sum(1 for r in rows if r['List'] == 'Primary')} primary / {sum(1 for r in rows if r['List'] == 'Backup')} backup, "
        f"{with_phone} with phone, {with_owner} with owner) -> {out / 'lsa_leads.xlsx'} + CSVs")


def write_workbook(path, rows, removed, kept):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    font = Font(name="Arial", size=10); bold = Font(name="Arial", size=10, bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7"); warn_fill = PatternFill("solid", fgColor="FDE9D9")
    widths = {"Business": 34, "Trade": 22, "City": 13, "Phone": 15, "Website": 30, "Owner": 20, "Title": 16, "Weekly hours": 40,
              "Why this lead ranks highly": 60, "Red flags": 45, "Notes": 45, "Highlights": 30, "Also listed for": 40,
              "LSA proof: profile URL": 30, "LSA proof: listing URL": 30, "Address (BBB)": 30, "Site signals": 26, "Owner source": 26}

    def sheet(title, data, cols):
        ws = wb.create_sheet(title[:31])
        ws.append(cols)
        for c in ws[1]:
            c.font = bold; c.fill = head_fill; c.alignment = Alignment(vertical="top", wrap_text=True)
        for r in data:
            ws.append([r.get(c, "") for c in cols])
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 12)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = font; c.alignment = Alignment(vertical="top", wrap_text=c.column_letter in "T U".split() or False)
        if "LSA display phone (DO NOT CALL - tracking number)" in cols:
            j = cols.index("LSA display phone (DO NOT CALL - tracking number)") + 1
            for row in ws.iter_rows(min_row=1, min_col=j, max_col=j):
                for c in row:
                    c.fill = warn_fill
        ws.freeze_panes = "F2" if "Business" in cols else "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    # README first
    ws = wb.active; ws.title = "README"
    lines = [
        ("Google Local Services Ads advertiser call list", bold),
        (f"Generated {now_iso()} by tools/lsa_leads/lsa_leads.py in the mowlid204.github.io repo.", font),
        ("", font),
        ("Tabs", bold),
        ("Top 50: the 50 with the strongest mix of LSA spend, missed-call exposure and reachable ownership.", font),
        ("Primary: ranks 1-300, the calling list. Caller 1 / 2 / 3: the Primary list split round-robin so nobody calls the same company.", font),
        ("Backup: ranks 301-500. All ranked: every eligible advertiser found, best to worst. Removed: advertisers dropped and why.", font),
        ("", font),
        ("Columns to use on a call: Phone (the company's own line), Owner, Title, Why this lead ranks highly, Red flags, Hours.", font),
        ("The column 'LSA display phone (DO NOT CALL - tracking number)' is the number inside the Google ad. Calling it bills the business for a lead. Evidence only.", bold),
        ("Verification: every row was captured from google.com/localservices/prolist (paying LSA advertisers only). Proof columns give the listing URL, profile URL, Google Ads customer ID, capture time and screenshot file in tools/lsa_leads/out/evidence/.", font),
        ("Phone source tells where the business line came from: BBB listing, the company website (tel: link or page text), or the BBB-built lists already in the repo.", font),
        ("Ranking: independent, established, 150-3,000 reviews, owner still involved, 24-hour or late phones, verified website and phone score highest. Corporate chains, private-equity-owned and acquired companies are removed; franchises and 5,000+ review operations are kept but penalised.", font),
    ]
    for i, (t, f) in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t).font = f
    ws.column_dimensions["A"].width = 140
    sheet("Top 50", [r for r in rows if r["Top 50"]], COLS)
    sheet("Primary", [r for r in rows if r["List"] == "Primary"], COLS)
    for c in (1, 2, 3):
        sheet(f"Caller {c}", [r for r in rows if r["Caller"] == c], COLS)
    sheet("Backup", [r for r in rows if r["List"] == "Backup"], COLS)
    sheet("All ranked", kept, COLS)
    sheet("Removed", removed, REMOVED_COLS)
    # summary with live counts
    ws = wb.create_sheet("Summary")
    ws.append(["Metric", "Count"]); ws["A1"].font = bold; ws["B1"].font = bold
    for i, (label, formula) in enumerate([
        ("Delivered (Primary + Backup)", "=COUNTA(Primary!E:E)+COUNTA(Backup!E:E)-2"),
        ("Primary", "=COUNTA(Primary!E:E)-1"), ("Backup", "=COUNTA(Backup!E:E)-1"),
        ("Caller 1", "=COUNTA('Caller 1'!E:E)-1"), ("Caller 2", "=COUNTA('Caller 2'!E:E)-1"),
        ("Caller 3", "=COUNTA('Caller 3'!E:E)-1"), ("Top 50", "=COUNTA('Top 50'!E:E)-1"),
        ("Eligible advertisers found", "=COUNTA('All ranked'!E:E)-1"), ("Removed", "=COUNTA(Removed!A:A)-1"),
        ("Primary rows with a phone", "=COUNTIF(Primary!I:I,\"?*\")"), ("Primary rows with an owner", "=COUNTIF(Primary!L:L,\"?*\")"),
    ], 2):
        ws.cell(row=i, column=1, value=label).font = font
        ws.cell(row=i, column=2, value=formula).font = font
    ws.column_dimensions["A"].width = 34
    wb.save(path)


def cmd_run(a):
    cmd_lists(a); cmd_profiles(a); cmd_enrich(a); cmd_export(a)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", choices=["lists", "profiles", "enrich", "fixsites", "export", "run"])
    ap.add_argument("--out", default=str(HERE / "out"))
    ap.add_argument("--cities", nargs="*", help="limit to these metro names, e.g. --cities Tulsa Omaha")
    ap.add_argument("--trades", nargs="*", help="limit to these trades: Roofing HVAC Plumbing Electrical")
    ap.add_argument("--limit", type=int, default=0, help="only the first N listing queries")
    ap.add_argument("--top", type=int, default=0, help="profiles/enrich: only the top N by score (0 = all); export: total delivered (default 500)")
    ap.add_argument("--primary", type=int, default=300, help="export: how many of the delivered rows form the primary calling list")
    ap.add_argument("--wait", type=float, default=5.0, help="seconds between Google page loads (plus jitter)")
    ap.add_argument("--headed", action="store_true", help="show the browser window")
    ap.add_argument("--no-retry", action="store_true", help="enrich: do not retry businesses that still have no phone")
    ap.add_argument("--shard", default="", help="enrich: write to enriched_<shard>.jsonl so several enrich runs can work in parallel")
    ap.add_argument("--no-search", action="store_true", help="enrich: skip the web search fallback for websites (BBB + domain guess only)")
    a = ap.parse_args(argv)
    if a.top == 0 and a.cmd in ("profiles", "enrich"):
        a.top = 10000
    Path(a.out).mkdir(parents=True, exist_ok=True)
    {"lists": cmd_lists, "profiles": cmd_profiles, "enrich": cmd_enrich, "fixsites": cmd_fixsites, "export": cmd_export, "run": cmd_run}[a.cmd](a)


if __name__ == "__main__":
    main()
